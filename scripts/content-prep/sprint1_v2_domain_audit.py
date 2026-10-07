#!/usr/bin/env python3
"""
DEMO-CONTENT-V2-PATCH-A — readiness audit of the non-Product domains of the
refreshed workbook (Homepage, Programme, Event, Talent, About). AUDIT ONLY:
nothing here is imported. Real people are reported as counts and row numbers,
never names or contact data. Drive links are reduced to "linked file / folder".

Usage: python3 scripts/content-prep/sprint1_v2_domain_audit.py <refreshed.xlsx> <out.json> --expect-sha256 <hex>
"""

import argparse
import collections
import hashlib
import json
import sys

import openpyxl


def url(cell):
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    return cell.value if isinstance(cell.value, str) and cell.value.startswith("http") else None


def kind(u):
    if not u:
        return None
    return "DRIVE_FOLDER" if "/folders/" in u else "DRIVE_FILE" if "/file/" in u else "OTHER"


def text(v):
    return " ".join(str(v).split()) if v not in (None, "") else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("out")
    ap.add_argument("--expect-sha256", required=True)
    a = ap.parse_args()
    raw = open(a.workbook, "rb").read()
    if hashlib.sha256(raw).hexdigest() != a.expect_sha256:
        sys.exit("source hash mismatch")
    import io

    wb = openpyxl.load_workbook(io.BytesIO(raw))

    hp = wb["Ảnh homepage"]
    programmes = []
    for r in range(3, 7):
        programmes.append({
            "row": r, "name": text(hp.cell(r, 3).value), "image": kind(url(hp.cell(r, 4))),
            "status_text": text(hp.cell(r, 6).value), "period": text(hp.cell(r, 8).value),
            "host": text(hp.cell(r, 10).value),
            "match": "PARTIAL" if url(hp.cell(r, 4)) else "MISSING",
        })
    selected = {x["name"] for x in json.load(open(a.out.rsplit("/", 1)[0] + "/02-products-selected-100.json"))}
    products = []
    for r in range(9, 20):
        name = text(hp.cell(r, 3).value)
        asset = kind(url(hp.cell(r, 4)))
        products.append({
            "row": r, "name": name, "logo": asset,
            "match": "READY" if asset == "DRIVE_FILE" and name in selected else "CONFLICT" if asset else "MISSING",
            "in_safe_product_set": name in selected,
        })
    news = [{"row": r, "image": kind(url(hp.cell(r, 3))), "match": "PARTIAL" if url(hp.cell(r, 3)) else "MISSING"} for r in range(22, 25)]

    lib = wb["Thư viện nhân tài"]
    rows = [r for r in range(3, lib.max_row + 1) if text(lib.cell(r, 3).value)]
    names = collections.Counter(text(lib.cell(r, 3).value) for r in rows)
    dup_rows = [[r for r in rows if text(lib.cell(r, 3).value) == n] for n, k in names.items() if k > 1]
    header = " ".join(str(c.value or "") for c in lib[1]).lower()
    det = wb["Nhân tài chi tiết"]
    drows = [r for r in range(3, det.max_row + 1) if text(det.cell(r, 4).value)]
    talent = {
        "rows": len(rows), "unique_names": len(names), "same_name_conflict_rows": dup_rows,
        "with_avatar": sum(1 for r in rows if url(lib.cell(r, 2))),
        "detail_rows": len(drows),
        "detail_rows_with_avatar_and_cover": sum(1 for r in drows if url(det.cell(r, 2)) and url(det.cell(r, 3))),
        "public_consent_field": any(k in header for k in ("consent", "đồng ý", "công khai")),
    }

    og = wb["Thành viên HUB_Orgchart"]
    orows = [r for r in range(2, og.max_row + 1) if text(og.cell(r, 3).value)]
    org = {
        "rows": len(orows), "with_photo_link": sum(1 for r in orows if url(og.cell(r, 2))),
        "groups": dict(sorted(collections.Counter(text(og.cell(r, 5).value) for r in orows).items())),
        "target_from_summary_sheet": text(wb["Tổng hợp"]["D12"].value),
    }
    summ = wb["Tổng hợp"]
    about = {
        "milestones_target": text(summ["D13"].value), "milestone_folder": kind(url(summ["F13"])),
        "hub_business_logos_target": text(summ["D14"].value), "hub_business_folder": kind(url(summ["F14"])),
        "hub_chapters_target": text(summ["D15"].value), "hub_chapter_folder": kind(url(summ["F15"])),
    }
    event = {"event_sheet_present": any("sự kiện" in ws.title.lower() for ws in wb.worksheets),
             "summary_target": text(summ["D6"].value)}
    out = {"SOURCE_FILE_HASH": a.expect_sha256, "homepage": {"programmes": programmes, "products": products, "news": news},
           "talent": talent, "org_chart": org, "about": about, "event": event}
    with open(a.out, "w", encoding="utf-8", newline="\n") as fh:
        fh.write(json.dumps(out, ensure_ascii=False, indent=2, sort_keys=True) + "\n")
    print(json.dumps(out, ensure_ascii=False, sort_keys=True)[:3000])


if __name__ == "__main__":
    main()
